"""
Document Ingestion Service
"""

import io
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.config import settings
from app.core.chunking import ChunkerFactory, TextChunk
from app.core.embeddings import get_embedding_pipeline
from app.core.retriever import get_retriever
from app.core.vector_store import get_vector_store
from app.models.document import DocumentChunk, DocumentType, IngestedDocument
from app.utils.helpers import chunk_list, generate_document_id, generate_uuid, get_file_extension

logger = logging.getLogger(__name__)


class IngestionService:
    SUPPORTED_TYPES = {".pdf": DocumentType.PDF, ".docx": DocumentType.DOCX, ".txt": DocumentType.TXT, ".md": DocumentType.MARKDOWN, ".html": DocumentType.HTML}

    def __init__(self, chunking_strategy: str = "recursive"):
        self._embedding_pipeline = get_embedding_pipeline()
        self._vector_store = get_vector_store()
        self._retriever = get_retriever()
        self._chunker = ChunkerFactory.create(chunking_strategy)

    async def ingest_file(self, file_content: bytes, filename: str, metadata: Optional[Dict[str, Any]] = None, collection_id: Optional[str] = None) -> IngestedDocument:
        extension = get_file_extension(filename)
        doc_type = self.SUPPORTED_TYPES.get(extension, DocumentType.UNKNOWN)
        if doc_type == DocumentType.UNKNOWN:
            raise ValueError(f"Unsupported file type '{extension}'.")
        text = await self._extract_text(file_content, doc_type)
        doc_id = generate_document_id(text, filename)
        doc = IngestedDocument(document_id=doc_id, title=Path(filename).stem, source=filename, document_type=doc_type, raw_content=text, metadata=metadata or {}, collection_id=collection_id)
        await self._process_document(doc)
        return doc

    async def ingest_text(self, title: str, content: str, source: Optional[str] = None, metadata: Optional[Dict[str, Any]] = None, collection_id: Optional[str] = None) -> IngestedDocument:
        doc_id = generate_document_id(content, source or title)
        doc = IngestedDocument(document_id=doc_id, title=title, source=source, document_type=DocumentType.TXT, raw_content=content, metadata=metadata or {}, collection_id=collection_id)
        await self._process_document(doc)
        return doc

    async def delete_document(self, document_id: str) -> int:
        deleted = self._vector_store.delete_by_document_id(document_id)
        self._vector_store.save()
        return deleted

    async def _process_document(self, doc: IngestedDocument) -> None:
        logger.info(f"Starting ingestion: document_id={doc.document_id}, title={doc.title}")
        raw_chunks: List[TextChunk] = self._chunker.chunk(doc.raw_content, metadata={"document_id": doc.document_id, "title": doc.title})
        doc_chunks: List[DocumentChunk] = []
        for tc in raw_chunks:
            chunk = DocumentChunk(
                chunk_id=generate_uuid(), document_id=doc.document_id, content=tc.content,
                chunk_index=tc.chunk_index, token_count=tc.token_count,
                metadata={**doc.metadata, "document_id": doc.document_id, "title": doc.title,
                           "source": doc.source, "chunk_index": tc.chunk_index,
                           "collection_id": doc.collection_id, "content": tc.content},
            )
            doc_chunks.append(chunk)
        doc.chunks = doc_chunks

        texts = [c.content for c in doc_chunks]
        all_embeddings: List[List[float]] = []
        for batch_texts in chunk_list(texts, settings.embedding_batch_size):
            batch_embs = await self._embedding_pipeline.embed_texts(batch_texts)
            all_embeddings.extend(batch_embs)

        for chunk, embedding in zip(doc_chunks, all_embeddings):
            chunk.embedding = embedding

        chunk_ids = [c.chunk_id for c in doc_chunks]
        embeddings = [c.embedding for c in doc_chunks]
        metadatas = [c.metadata for c in doc_chunks]
        self._vector_store.add_embeddings(chunk_ids, embeddings, metadatas)
        self._vector_store.save()
        self._retriever.index_chunks(chunk_ids, [c.content for c in doc_chunks], metadatas)
        logger.info(f"Ingestion complete: document_id={doc.document_id}, chunks={len(doc_chunks)}")

    async def _extract_text(self, content: bytes, doc_type: DocumentType) -> str:
        if doc_type == DocumentType.PDF:
            return self._extract_pdf(content)
        elif doc_type == DocumentType.DOCX:
            return self._extract_docx(content)
        elif doc_type in (DocumentType.TXT, DocumentType.MARKDOWN):
            return content.decode("utf-8", errors="replace")
        elif doc_type == DocumentType.HTML:
            return self._extract_html(content)
        raise ValueError(f"No extractor for {doc_type}")

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            raise ImportError("pypdf is required")

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            raise ImportError("python-docx is required")

    @staticmethod
    def _extract_html(content: bytes) -> str:
        import re
        text = content.decode("utf-8", errors="replace")
        return re.sub(r"<[^>]+>", " ", text)
